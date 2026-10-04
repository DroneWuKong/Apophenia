#!/usr/bin/env python3
"""Generate the maintained Apophenia user guide in DOCX and standalone HTML."""

from __future__ import annotations

import base64
import html
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
IMAGES = DOCS / "images"
DOCX_OUT = DOCS / "Apophenia_User_Guide.docx"
HTML_OUT = DOCS / "Apophenia_User_Guide.html"

INK = "111318"
MUTED = "535A69"
LINE = "D9D9D9"
HEADER = "222B3D"
ALT_ROW = "F1F4FA"
LAVENDER = "A9BFFD"
RED = "B42318"


def set_font(run, size=11, color=INK, bold=False, name="Aptos") -> None:
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


def add_run(paragraph, text: str, *, size=11, color=INK, bold=False, italic=False):
    run = paragraph.add_run(text)
    set_font(run, size=size, color=color, bold=bold)
    run.italic = italic
    return run


def set_cell_fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=130, start=150, bottom=130, end=150) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "5")
        node.set(qn("w:color"), LINE)


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def configure_doc() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.14

    title = doc.styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string("000000")
    title.paragraph_format.space_after = Pt(10)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for name, size in (("Heading 1", 21), ("Heading 2", 14)):
        style = doc.styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string("000000")
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)

    doc.core_properties.title = "Apophenia Total Capture User Guide"
    doc.core_properties.subject = "Capture gates, local evidence, export formats, and analysis boundaries"
    doc.core_properties.author = "DroneWuKong"
    return doc


def add_body(doc: Document, text: str, *, bold_lead: str | None = None) -> None:
    p = doc.add_paragraph()
    if bold_lead:
        add_run(p, bold_lead, bold=True)
    add_run(p, text)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        add_run(p, item)


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        add_run(p, item)


def add_figure(doc: Document, filename: str, caption: str, alt: str, width: float) -> None:
    path = IMAGES / filename
    if not path.is_file():
        raise FileNotFoundError(path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    inline = p.add_run().add_picture(str(path), width=Inches(width))
    inline._inline.docPr.set("descr", alt)
    caption_p = doc.add_paragraph()
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p.paragraph_format.space_after = Pt(9)
    add_run(caption_p, caption, size=9, color=MUTED, italic=True)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_header(header)
    for index, (cell, label) in enumerate(zip(header.cells, headers)):
        cell.width = Inches(widths[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_fill(cell, HEADER)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_run(p, label, size=10.5, color="FFFFFF", bold=True)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, (cell, value) in enumerate(zip(cells, values)):
            cell.width = Inches(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index % 2 == 1:
                set_cell_fill(cell, ALT_ROW)
            p = cell.paragraphs[0]
            add_run(p, value, size=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def page_break(doc: Document) -> None:
    doc.add_page_break()


def build_docx() -> None:
    doc = configure_doc()

    p = doc.add_paragraph(style="Title")
    add_run(p, "Apophenia Total Capture User Guide", size=30, color="000000", bold=True)
    p = doc.add_paragraph()
    add_run(p, "What the app records, how authorization works, what stays local, and how to export for people and machines", size=14, color=MUTED)
    add_body(
        doc,
        "Apophenia is a personal Android and Garmin black box for reconstructing the circumstances around a timestamped event. The operator logs the moment; the software freezes only the context that was deliberately authorized."
    )
    add_body(
        doc,
        "It can be maximum invasive when everything is authorized and armed. That is the capability ceiling, not the default. There is no account, analytics SDK, automatic cloud sync, or background uploader. Data stays local until an export is built, verified, previewed, and deliberately routed.",
        bold_lead="The privacy bargain. "
    )
    add_figure(
        doc,
        "apophenia-home-v03.png",
        "Current debug application on an Android emulator. The event timestamp is saved before context enrichment.",
        "Apophenia capture screen with THAT WAS WEIRD, five VIBE choices, egress, and the live status strip",
        2.1,
    )

    page_break(doc)
    doc.add_heading("What the app is", level=1)
    add_body(
        doc,
        "The instrument is designed to answer a bounded question: what phone, environmental, radio, physiological, vehicle, aircraft, and audiovisual circumstances were present around the observation? It stores observations as descriptions and keeps hypotheses and evaluations separate."
    )
    add_body(doc, "Four controls must not be conflated:")
    add_numbered(doc, [
        "The named Apophenia gate authorizes a capability.",
        "Android permission or MediaProjection consent is granted separately where required.",
        "A drive, flight, control-link, RF, microphone, camera, or screen session is started explicitly where the channel has a live state.",
        "An export remains local through preparation and manifest preview; only an explicit route crosses the app boundary.",
    ])
    add_body(
        doc,
        "Gate presence or simulator success does not prove that hardware exists, that Android exposes the channel, that a recording is live, or that a physical result has been validated.",
        bold_lead="Evidence boundary. "
    )

    page_break(doc)
    doc.add_heading("Capture a moment", level=1)
    add_figure(
        doc,
        "apophenia-home-annotated-v03.png",
        "The source UI is an emulator capture; arrows and labels are documentation overlays.",
        "Annotated capture screen with arrows to live status, timestamp-first logging, and the distinct egress event",
        4.55,
    )
    add_bullets(doc, [
        "THAT WAS WEIRD saves the tap time first and enriches it with authorized context afterward.",
        "VIBE records a one-tap ordinal state from 1 through 5; holding opens an optional note without moving the event time.",
        "NOPE, I'M OUT records VIBE 5 with egress=true, preserving the difference between leaving and staying.",
        "The status strip distinguishes authorized gate count from AV rings that are actually live.",
    ])

    page_break(doc)
    doc.add_heading("Choose how invasive the session should be", level=1)
    add_figure(
        doc,
        "apophenia-settings-annotated-v03.png",
        "TOTAL_EVIDENCE and presets change named gate authorization. Permissions and hardware starts remain separate.",
        "Annotated Settings screen with arrows to TOTAL_EVIDENCE, the permission boundary, and the FIELD preset",
        6.7,
    )
    add_table(doc, ["Control", "Does", "Does not"], [
        ["Named gate", "Authorizes one channel", "Bypass Android or invent hardware"],
        ["TOTAL_EVIDENCE", "Arms all capture gates and maximum configured buffers", "Start OBD UAS RF or screen capture or enable LAN export"],
        ["FIELD", "Arms flight link detector RF ground AV and watch gates", "Claim aircraft radio SDR camera or watch validation"],
        ["DRIVE", "Arms vehicle cabin presence and AV gates", "Connect an adapter or start a drive session"],
        ["EVERYTHING", "Arms every capture gate", "Send any data or silently grant permissions"],
    ], [1.25, 2.55, 2.65])

    page_break(doc)
    doc.add_heading("What it can capture", level=1)
    add_body(doc, "Every channel is visible even when unavailable. Omniprobe records the stored value or a gap reason: gate off, permission denied, platform restricted, hardware absent, or statute.")
    add_table(doc, ["Domain", "Examples", "Activation boundary"], [
        ["Phone and environment", "Sensors battery thermal display network Wi-Fi Bluetooth time weather space weather", "Named gates plus platform permissions where required"],
        ["Protected contents", "Notifications calendar contacts message metadata", "Tier 2 exact-name authorization and Android access"],
        ["Physiology", "Health Connect and Garmin context", "Read permission or paired-phone delivery"],
        ["Vehicle", "OBD-II Automotive EV and cabin Bluetooth", "Gate paired adapter and operator-started DRIVE_SESSION"],
        ["Aircraft and field", "MAVLink CRSF GHST Field-Kit TAK ground context and SDR", "Gate owned hardware or selected transport and bounded session/window"],
        ["Audio and video", "Microphone main front concurrent cameras and screen", "Tier 2 gate permission or per-arm consent and persistent live indicator"],
    ], [1.35, 2.8, 2.3])
    doc.add_heading("Local retention", level=2)
    add_body(doc, "Ordinary records and derived metrics live in SQLite. Tier 2 contents and raw AV are encrypted with device-local keys. Raw AV expires after the configured deadline, 14 days by default, unless that event is marked keep forever.")
    add_bullets(doc, [
        "Scrub and expiry delete media ciphertext, its sidecar, and its event key.",
        "Non-reconstructive loudness, spectral, motion, brightness, flicker, and scene-change metrics survive.",
        "A purge ledger records what was deleted and why.",
        "Raw BSSIDs, MAC addresses, adapter addresses, and airframe IDs are not durable fields; locally keyed hashes are stored instead.",
    ])

    page_break(doc)
    doc.add_heading("Export for people and machines", level=1)
    add_figure(
        doc,
        "apophenia-settings-export-v03.png",
        "All export routes prepare locally and show a verified manifest before a destination can be selected.",
        "Settings export actions for data-only full evidence raw SQLite backup and restore",
        3.85,
    )
    doc.add_heading("Choose an export format", level=2)
    add_table(doc, ["Export", "Human-readable", "Machine-readable"], [
        ["Data-only", "Analysis README and CSV tables", "Canonical JSON CSV and JSON data dictionary"],
        ["Raw SQLite", "Schema documentation", "Checkpointed integrity-checked database"],
        ["Event dossier", "Plain summary and SVG chart", "Event JSON context CSV inventories and retained evidence"],
        ["Selected report", "Self-contained HTML and PDF", "Report JSON context CSV and SVG"],
        ["Full evidence", "Analysis README and inventories", "JSON CSV AV Tier 2 attachments RF and hashes"],
    ], [1.35, 2.35, 2.75])

    page_break(doc)
    doc.add_heading("Manifest preview and release", level=1)
    add_figure(
        doc,
        "apophenia-manifest-v03.png",
        "The preview lists payload count and bytes, AV and Tier 2 flags, the bundle digest, and every declared file digest.",
        "Data-only manifest preview with hashes and explicit raw AV and Tier 2 exclusions",
        3.95,
    )
    add_body(doc, "The default data-only tier excludes raw AV, Tier 2 plaintext, RF IQ, and inbound attachment bytes. Full evidence materializes device-bound evidence into portable plaintext and requires two confirmations.")
    add_body(doc, "Sharesheet handoff, completed SAF writes, and LAN endpoint acknowledgement are different audit outcomes. None proves recipient reading, durable retention, or exactly-once delivery.", bold_lead="Honest route receipts. ")

    doc.add_heading("Analysis boundaries", level=1)
    add_bullets(doc, [
        "Group dense rows by capture_id before comparison; rows in one capture are not independent observations.",
        "Exclude phase POST from predictors while retaining it for reconstruction.",
        "Keep is_control and CONTROL phase rows identifiable.",
        "Preserve the tested-channel count and multiple-comparisons correction.",
        "Do not convert association into causation diagnosis or explanation.",
        "Treat a refuted preregistered pattern as the instrument working.",
    ])
    add_body(doc, "Software tests and emulator screenshots do not establish phone hardware, multicamera, microphone, radio, adapter, vehicle, aircraft, SDR, watch, field, or flight performance. Use the dated physical-validation checklist in PROJECT_HANDOFF.md.", bold_lead="Validation boundary. ")

    DOCX_OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX_OUT)


def data_uri(filename: str) -> str:
    path = IMAGES / filename
    encoded = base64.b64encode(path.read_bytes()).decode("ascii")
    return f"data:image/png;base64,{encoded}"


def build_html() -> None:
    capture = data_uri("apophenia-home-annotated-v03.png")
    settings = data_uri("apophenia-settings-annotated-v03.png")
    export = data_uri("apophenia-settings-export-v03.png")
    manifest = data_uri("apophenia-manifest-v03.png")
    rows = [
        ("Data-only", "README and CSV", "Canonical JSON, CSV, data dictionary"),
        ("Raw SQLite", "Schema documentation", "Checkpointed .db"),
        ("Event dossier", "Summary and SVG", "JSON, CSV, inventories, evidence"),
        ("Selected report", "HTML and PDF", "JSON, CSV, SVG"),
        ("Full evidence", "README and inventories", "JSON/CSV plus AV, Tier 2, attachments, RF"),
    ]
    table_rows = "".join(
        f"<tr><td>{html.escape(a)}</td><td>{html.escape(b)}</td><td>{html.escape(c)}</td></tr>"
        for a, b, c in rows
    )
    page = f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Apophenia Total Capture User Guide</title>
<style>
:root{{--bg:#0b0e14;--panel:#141a25;--ink:#f4f6fb;--muted:#b8c0d2;--line:#303a4f;--lav:#a9bffd;--red:#ff6b61}}
*{{box-sizing:border-box}}body{{margin:0;background:var(--bg);color:var(--ink);font:17px/1.62 system-ui,sans-serif}}
main{{width:min(1080px,calc(100% - 32px));margin:auto;padding:48px 0 96px}}h1{{font-size:clamp(42px,7vw,78px);line-height:1;letter-spacing:-.045em;margin:0 0 20px}}h2{{font-size:32px;margin:56px 0 14px}}h3{{font-size:22px}}
p,li{{max-width:850px}}.lead{{font-size:22px;color:var(--muted)}}strong{{color:#fff}}.boundary{{color:var(--lav);font-weight:700}}
figure{{margin:28px 0;background:var(--panel);border:1px solid var(--line)}}figure img{{display:block;width:100%;height:auto}}figcaption{{padding:12px 16px;color:var(--muted);font-size:13px}}
.phone{{max-width:520px}}table{{width:100%;border-collapse:collapse;margin:24px 0}}th,td{{border:1px solid var(--line);padding:13px;text-align:left;vertical-align:top}}th{{background:#222b3d}}tr:nth-child(even){{background:#111722}}
.grid{{display:grid;grid-template-columns:1fr 1fr;gap:22px}}@media(max-width:800px){{.grid{{grid-template-columns:1fr}}}}
</style></head><body><main>
<h1>Apophenia Total Capture User Guide</h1>
<p class="lead">What the app records, how authorization works, what stays local, and how to export for people and machines.</p>
<p>Apophenia is a personal Android and Garmin black box for reconstructing the circumstances around a timestamped event. It can be <strong>maximum invasive</strong> when everything is deliberately authorized and armed. That is the capability ceiling, not the default.</p>
<p><span class="boundary">The privacy bargain:</span> no account, analytics SDK, automatic cloud sync, or background uploader. Data stays local until an export is built, verified, previewed, and deliberately routed.</p>
<h2>Capture the moment</h2><figure><img alt="Annotated capture screen" src="{capture}"><figcaption>Timestamp first, authorized context second. Egress is a distinct event class.</figcaption></figure>
<h2>Choose how invasive the session should be</h2><figure><img alt="Annotated TOTAL EVIDENCE and preset settings" src="{settings}"><figcaption>Gates, Android permissions, and live hardware starts remain separate.</figcaption></figure>
<p>TOTAL_EVIDENCE and presets arm named capture gates. They do not grant OS permissions, accept screen-record consent, start OBD/UAS/RF sessions, or open an export route.</p>
<h2>What stays local</h2><p>Ordinary evidence and derived metrics live in SQLite. Tier 2 contents and retained AV are encrypted with device-local keys. Raw media is bounded and expires on its retention schedule unless an event is marked keep forever. Purge receipts remain after media is scrubbed.</p>
<h2>Export for people and machines</h2><div class="grid"><figure class="phone"><img alt="Export choices" src="{export}"><figcaption>Prepare locally before choosing a route.</figcaption></figure><figure class="phone"><img alt="Manifest preview" src="{manifest}"><figcaption>Review files, sizes, hashes, and sensitive-content flags.</figcaption></figure></div>
<table><thead><tr><th>Export</th><th>Human-readable</th><th>Machine-readable</th></tr></thead><tbody>{table_rows}</tbody></table>
<h2>Analysis boundaries</h2><ul><li>Group dense rows by <code>capture_id</code>.</li><li>Exclude <code>phase=POST</code> from predictors.</li><li>Keep control windows identifiable.</li><li>Preserve multiple-comparisons correction.</li><li>Correlation is not causation.</li><li>A refuted preregistration is a useful result.</li></ul>
<p><span class="boundary">Validation boundary:</span> emulator and simulator evidence proves software behavior only, not physical sensors, AV hardware, radios, adapters, vehicles, aircraft, SDRs, watches, field performance, or flight performance.</p>
</main></body></html>"""
    HTML_OUT.write_text(page, encoding="utf-8")


if __name__ == "__main__":
    build_docx()
    build_html()
    print(DOCX_OUT)
    print(HTML_OUT)
